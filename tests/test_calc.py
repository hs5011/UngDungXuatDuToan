import io

import openpyxl
import pytest

from dutoan import scaling
from dutoan.calc import compute, compute_fees, compute_pm, _compute_at, kdp_max, rnd
from dutoan.catalog import TASK_BY_ID
from dutoan.defaults import default_project
from dutoan.excel_export import export_bytes


@pytest.fixture
def proj():
    p = default_project()
    p["tham_so"]["kdp"] = 0.05
    return p


@pytest.mark.parametrize("gms, expect", [(1e9, (1, 6)), (2e9, (1, 11)), (3e9, (2, 8)), (6e9, (2, 15)), (10e9, (3, 16))])
def test_vi_du_phan_tich_nghiep_vu(proj, gms, expect):
    """Ví dụ đã duyệt (YEU-CAU-FINAL mục 8): việc 'Phân tích nghiệp vụ, đặc tả UC'."""
    ts = proj["tham_so"]
    total = scaling.total_days(TASK_BY_ID["B2"].anchors, gms, ts)
    assert scaling.split(total, ts["bac_thang"]) == expect


def test_viec_xuat_hien_theo_nguong(proj):
    ts = proj["tham_so"]
    assert not scaling.is_active(TASK_BY_ID["A6"].anchors, 0.9e9, ts)
    assert scaling.is_active(TASK_BY_ID["A6"].anchors, 1e9, ts)
    assert not scaling.is_active(TASK_BY_ID["A8"].anchors, 4.9e9, ts)
    assert scaling.is_active(TASK_BY_ID["A8"].anchors, 5e9, ts)


def test_qlda_moc_S_khop_49_trieu(proj):
    """Mốc neo A tại quy mô S: ≈ 49 tr sau VAT (Cql 55%, TN 6%, VAT 8%)."""
    proj["tham_so"]["gms"] = 0.999_999e9
    g = compute(proj).groups["A"]
    assert [(r.experts, r.days) for r in g.rows] == [(1, 2), (1, 8), (1, 4), (1, 2), (1, 4)]
    assert g.total == pytest.approx(49_138_346, abs=1)


def test_cql_theo_bac_thang(proj):
    from dutoan.calc import cql_rate
    ts = proj["tham_so"]
    assert cql_rate(1e9, ts) == 0.55
    assert cql_rate(1e9 + 1, ts) == 0.50
    assert cql_rate(5e9, ts) == 0.45


def test_pm_noi_bo_suy_nguoc_dung_gms(proj):
    pm = compute_pm(3e9, proj["tham_so"])
    assert pm["total"] == pytest.approx(3e9)
    assert sum(r["thanh_tien"] for r in pm["roles"]) == pytest.approx(pm["nc"])
    assert sum(r["ty_le"] for r in pm["roles"]) == pytest.approx(1.0)


def test_don_gia_ky_su_khop_template(proj):
    """Dongia_NgayCong: KS-B1 = 2,34 × 2.530.000 × (1 + 1,2) + 2,34 × 2.530.000 × 21,5%."""
    from dutoan.calc import ks_month_price
    mlcb = 2.34 * 2_530_000
    assert ks_month_price(1, proj["tham_so"]) == pytest.approx(mlcb * 2.2 + mlcb * 0.215)


def test_che_do_gms_du_phong_khong_vuot_toi_da(proj):
    for gms in (0.5e9, 3e9, 12e9):
        proj["tham_so"]["gms"] = gms
        r = compute(proj)
        assert r.tmdt % 1_000_000 == 0
        assert 0 <= r.totals["V"] <= kdp_max(proj["tham_so"]) * r.totals["base"] + 1
        assert r.tmdt == pytest.approx(sum(i.post for i in r.items))


def test_che_do_tmdt_gms_lon_nhat(proj):
    ts = proj["tham_so"]
    ts["nhap"], ts["tmdt"], ts["du_phong_toi_thieu"] = "TMDT", 5e9, 20e6
    r = compute(proj)
    assert r.tmdt == 5e9
    assert r.totals["V"] >= 20e6
    assert _compute_at(proj, r.gms + 1000, 5e9).totals["V"] < 20e6


def test_thue_pm_bang_gms_chia_khau_hao(proj):
    ts = proj["tham_so"]
    ts["che_do"], ts["so_nam_thue"], ts["gms"] = "THUE", 3, 2e9
    r = compute(proj)
    assert r.item("thue_pm").post == pytest.approx(2e9 / 8 * 3)
    assert sum(r.phan_bo["years"]) == pytest.approx(r.tmdt)


def test_kiem_toan_tham_tra_toi_thieu(proj):
    f = compute_fees(1e6, 1e6, 0, proj["tham_so"])
    assert f["kiem_toan_pre"] == 1_000_000
    assert f["kiem_toan_vat"] == 80_000
    assert f["tham_tra"] == 500_000


def test_noi_suy_nd193(proj):
    """TMĐT 7,5 tỷ: kiểm toán nội suy giữa 0,96% (5 tỷ) và 0,645% (10 tỷ)."""
    f = compute_fees(7.5e9, 7e9, 0, proj["tham_so"])
    assert f["ty_le_kt"] == pytest.approx(rnd(0.96 - (0.96 - 0.645) / 5 * 2.5, 4))


def test_nd214_dung_van_ban(proj):
    proj["tham_so"]["gms"] = 3e9
    proj["ghi_de"] = {"E2": {"dung_vb": True}}
    r = compute(proj)
    e2 = next(x for x in r.groups["E"].rows if x.id == "E2")
    gia_goi = sum(i.pre for i in r.items if i.sec == "I") / r.counts["goi"]
    assert e2.amount == 0
    assert e2.ref_pre == rnd(min(max(gia_goi * 0.002, 3e6), 60e6)) * r.counts["goi"]


def test_ghi_de_so_chuyen_gia(proj):
    proj["ghi_de"] = {"B2": {"so_cg": 4}}
    r = compute(proj)
    b2 = next(x for x in r.groups["B"].rows if x.id == "B2")
    assert b2.experts == 4 and b2.overridden


def test_tham_tra_thiet_ke_tu_dong_tat_khi_co_pm(proj):
    ids = [x.id for x in compute(proj).groups["C"].rows]
    assert "C3" not in ids
    proj["tham_so"]["co_pm_noi_bo"] = False
    ids = [x.id for x in compute(proj).groups["C"].rows]
    assert "C3" in ids


def test_xuat_excel_co_cong_thuc_song(proj):
    data = export_bytes(proj, compute(proj))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "BẢNG TMĐT" in wb.sheetnames and "3_GMS_PM_NOI_BO" in wb.sheetnames
    for name in ("VAT", "GMS", "TMDT", "KDP", "DG_M3", "KS_TB_THANG", "GIA_GOI"):
        assert name in wb.defined_names
    assert wb.calculation.iterate
    ws = wb["BẢNG TMĐT"]
    assert any(isinstance(c.value, str) and c.value.startswith("=") for row in ws.iter_rows() for c in row)


def test_cong_viec_bo_sung_va_xoa(proj):
    proj["tham_so"]["gms"] = 3e9
    before = compute(proj).groups["B"]
    proj["viec_them"] = {"B": [{"ten": "Khảo sát bổ sung", "bac": "M2", "so_cg": 2, "so_ngay": 3, "he_so": 1}]}
    proj["ghi_de"] = {"B1": {"bo": True}}
    g = compute(proj).groups["B"]
    them = next(r for r in g.rows if r.id == "B+1")
    assert them.amount == pytest.approx(2 * 3 * rnd(55_000_000 / 26))
    b1 = next(r for r in g.rows if r.id == "B1")
    assert b1.excluded and b1.amount == 0
    b1_old = next(r for r in before.rows if r.id == "B1").amount
    assert g.ccg == pytest.approx(before.ccg - b1_old + them.amount)


def test_doi_ten_viec_mac_dinh_va_ghi_chu_mac_dinh(proj):
    proj["ghi_de"] = {"A1": {"ten": "Lập kế hoạch tổng thể (đã sửa)"}}
    r = compute(proj)
    a1 = next(x for x in r.groups["A"].rows if x.id == "A1")
    assert a1.name == "Lập kế hoạch tổng thể (đã sửa)" and a1.overridden
    assert all(i.note for i in r.items)


def test_so_sanh_van_ban_co_cau_truc(proj):
    r = compute(proj)
    ma = {c["ma"] for c in r.compare}
    assert {"E2", "E3", "F2", "F3"} <= ma
    assert all(c["dang_dung"] == "Chuyên gia" for c in r.compare)
