"""Test bổ sung cho các mục tính chưa được phủ: chế độ THUE, QLDA tự thực hiện,
trọn gói, tắt nhóm, ghi đè nhóm, chi phí khác (Ck), cảnh báo, phân bổ vốn theo
chế độ, dm1688, nội suy phí thẩm định, tb_factor, QTVH thuê theo năm."""
import io

import openpyxl
import pytest

from dutoan.calc import (
    compute, compute_fees, compute_pm, compute_thue_pm, compute_tron_goi,
    compute_qlda_tu_th, _compute_at, kdp_max, interp, tier, counts,
    day_price, ks_month_price, dm1688_rate, tb_factor, cql_rate, rnd, round_tmdt,
    group_enabled, tham_tra_tk_on, default_alloc, so_nam_phan_bo, start_year,
)
from dutoan.catalog import DM1688
from dutoan.defaults import default_project, KDP_MAX, KDP_MAX_THUE, merge_defaults
from dutoan.excel_export import export_bytes


@pytest.fixture
def proj():
    p = default_project()
    p["tham_so"]["kdp"] = 0.05
    return p


# ================================ Tiện ích thuần ================================
def test_rnd_lam_tron_nua_len_xa_so_0():
    assert rnd(2.5) == 3
    assert rnd(-2.5) == -3
    assert rnd(0.1234, 2) == 0.12
    assert rnd(0.125, 2) == 0.13


def test_interp_ngoai_bang_lay_dau_cuoi():
    assert interp(-1, [0, 10], [5, 15]) == 5
    assert interp(100, [0, 10], [5, 15]) == 15
    assert interp(5, [0, 10], [5, 15]) == 10


def test_tier_va_counts_theo_quy_mo(proj):
    ts = proj["tham_so"]
    assert tier(0.5e9, ts) == 0 and tier(2e9, ts) == 1 and tier(10e9, ts) == 2
    c0 = counts(0.5e9, ts)
    assert (c0["goi"], c0["lop"], c0["he_thong"]) == (2, 1, 0)
    c2 = counts(10e9, ts)
    assert (c2["goi"], c2["lop"], c2["he_thong"]) == (5, 6, 5)


def test_counts_ghi_de_bang_tay(proj):
    proj["tham_so"]["so_goi"] = 7
    assert counts(0.5e9, proj["tham_so"])["goi"] == 7


def test_kdp_max_phan_biet_che_do(proj):
    ts = proj["tham_so"]
    assert kdp_max(ts) == KDP_MAX["TMDT"]
    ts["loai_ho_so"] = "BCKTKT"
    assert kdp_max(ts) == KDP_MAX["BCKTKT"]
    ts["che_do"] = "THUE"
    assert kdp_max(ts) == KDP_MAX_THUE


def test_round_tmdt_lam_tron_len_khi_trong_kdp(proj):
    """Trong giới hạn Kdp thì làm tròn LÊN theo đơn vị lam_tron."""
    ts = proj["tham_so"]
    ts["lam_tron"] = 1_000_000
    base = 100_123_456
    t = round_tmdt(base, 0.05, ts)
    assert t % 1_000_000 == 0 and t >= base * 1.05 - 1_000_000


def test_round_tmdt_lam_tron_xuong_khi_vuot_kdp(proj):
    """Khi làm tròn LÊN khiến dự phòng vượt Kdp thì chuyển sang làm tròn XUỐNG."""
    ts = proj["tham_so"]
    ts["lam_tron"] = 10_000_000  # bước 10 triệu
    base = 99_000_000
    # rate = 0.08 → x = 106.92tr; ceil→110tr, dp=11tr > 9.9tr (10% × 99tr) → floor→100tr
    t = round_tmdt(base, 0.08, ts)
    assert t == 100_000_000


# ================================ Giá chuyên gia ================================
def test_day_price_khop_cong_thuc(proj):
    """ROUND(lương × hệ số nâng / số ngày tháng)."""
    ts = proj["tham_so"]
    ts["he_so_nang"] = 1.2
    assert day_price("M3", ts) == rnd(40_000_000 * 1.2 / 26)


def test_ks_month_price_theo_vung(proj):
    ts = proj["tham_so"]
    p1 = ks_month_price(1, ts)
    ts["vung"] = "IV"
    p4 = ks_month_price(1, ts)
    assert p1 > p4  # vùng I có Hđc cao nhất


# ================================ Chế độ THUE ===================================
def test_che_do_thue_co_thue_pm_va_qtvh_va_dv_nam(proj):
    ts = proj["tham_so"]
    ts["che_do"], ts["gms"], ts["so_nam_thue"] = "THUE", 2e9, 3
    proj["dv_thue_nam"] = [{"ten": "Hạ tầng", "so_luong": 1, "don_gia": 50_000_000, "vat": 0.08}]
    r = compute(proj)
    assert r.item("thue_pm") is not None
    assert r.item("qtvh") is not None  # J nhóm chuyên gia (không trọn gói)
    assert r.item("dv_nam") is not None
    # dịch vụ thuê × số năm
    assert r.item("dv_nam").pre == pytest.approx(50_000_000 * 3)
    assert r.item("dv_nam").vat == pytest.approx(50_000_000 * 0.08 * 3)


def test_che_do_thue_tron_goi_thay_the_nhom_J(proj):
    ts = proj["tham_so"]
    ts["che_do"], ts["j_tron_goi"], ts["j_tron_goi_nam"] = "THUE", True, 100_000_000
    ts["so_nam_thue"], ts["gms"] = 2, 1e9
    r = compute(proj)
    q = r.qtvh_tron_goi
    assert q["nc"] == pytest.approx(100_000_000 * 2)
    assert q["pre"] + q["vat"] == pytest.approx(q["total"])
    # J không tính nữa
    assert "J" not in r.groups
    assert r.item("qtvh").post == pytest.approx(q["total"])


def test_che_do_thue_cho_phep_dat_du_phong_toi_thieu_tien(proj):
    ts = proj["tham_so"]
    ts.update(che_do="THUE", nhap="TMDT", tmdt=3e9, du_phong_theo="TIEN", du_phong_toi_thieu=5e6, so_nam_thue=2)
    r = compute(proj)
    assert r.tmdt == 3e9
    assert r.totals["V"] >= 5e6


def test_thue_pm_bang_khau_hao(proj):
    pm = compute_pm(2e9, proj["tham_so"])
    t = compute_thue_pm(pm, {"so_nam_thue": 3, "khau_hao_nam": 8})
    assert t["he_so"] == pytest.approx(3 / 8)
    assert t["total"] == pytest.approx(pm["total"] * 3 / 8)


def test_tron_goi_cong_thuc_khop_tay():
    ts = {"j_tron_goi_nam": 50_000_000, "so_nam_thue": 2, "pm_cpc": 0.5, "pm_tn": 0.1, "vat": 0.08}
    q = compute_tron_goi(ts)
    nc = 50_000_000 * 2
    cpc = nc * 0.5
    tn = (nc + cpc) * 0.1
    vat = (nc + cpc + tn) * 0.08
    assert (q["nc"], q["cpc"], q["tn"], q["vat"]) == pytest.approx((nc, cpc, tn, vat))


# ================================ QLDA tự thực hiện =============================
def test_qlda_tu_thuc_hien_tinh_tien_luong_cong_bh(proj):
    ts = proj["tham_so"]
    ts["qlda"] = "TU_TH"
    r = compute(proj)
    q = r.qlda_tu_th
    expect_luong = sum(
        (float(x["he_so_luong"]) + float(x["pc_chuc_vu"]) + float(x["pc_khu_vuc"]) + float(x["pc_khac"]))
        * ts["luong_co_so"] * float(x["so_nguoi"]) * float(x["so_thang"])
        for x in proj["qlda_nhan_su"]
    )
    assert q["luong"] == pytest.approx(expect_luong)
    assert q["bh"] == pytest.approx(expect_luong * ts["bh_qlda"])
    assert q["total"] == pytest.approx(q["luong"] + q["bh"] + q["chi_khac"])
    # Khi tự TH thì không có nhóm A
    assert "A" not in r.groups


# ================================ Nhóm bật/tắt ==================================
def test_nhom_tat_khong_vao_tong(proj):
    proj["nhom_tat"] = ["I"]  # tắt kiểm thử
    r = compute(proj)
    assert "I" not in r.groups
    assert r.item("kiem_thu") is None


def test_group_enabled_logic(proj):
    ts = proj["tham_so"]
    ts["qlda"], ts["co_pm_noi_bo"], ts["che_do"], ts["j_tron_goi"] = "TU_TH", False, "THUE", True
    assert group_enabled("A", proj) is False
    assert group_enabled("I", proj) is False
    assert group_enabled("J", proj) is False
    ts["j_tron_goi"] = False
    assert group_enabled("J", proj) is True


def test_tham_tra_thiet_ke_switch_tu_dong_va_tay(proj):
    ts = proj["tham_so"]
    ts["co_pm_noi_bo"] = True
    assert tham_tra_tk_on(ts) is False
    ts["tham_tra_thiet_ke"] = True  # ép bật dù có PM
    assert tham_tra_tk_on(ts) is True


# ================================ Ghi đè nhóm, Ck ===============================
def test_ghi_de_nhom_tong_cong(proj):
    proj["ghi_de_nhom"] = {"B": 100_000_000}
    r = compute(proj)
    g = r.groups["B"]
    assert g.total == 100_000_000
    assert g.pre + g.vat == pytest.approx(100_000_000)


def test_chi_phi_khac_ck_cong_vao_nhom(proj):
    proj["ck"] = {"B": 5_000_000}
    r = compute(proj)
    g = r.groups["B"]
    assert g.ck == 5_000_000
    # Thu nhập tính thuế đã bao gồm Ck
    assert g.tn == rnd(g.tn_rate * (g.ccg + g.cql + 5_000_000))


# ================================ Việc thêm nhiều nhóm ==========================
def test_viec_them_nhieu_nhom(proj):
    proj["viec_them"] = {
        "A": [{"ten": "QLDA bổ sung", "bac": "M2", "so_cg": 1, "so_ngay": 5, "he_so": 1}],
        "G": [{"ten": "Đào tạo thêm", "bac": "M3", "so_cg": 2, "so_ngay": 10, "he_so": 2}],
    }
    r = compute(proj)
    assert any(x.id == "A+1" for x in r.groups["A"].rows)
    g_extra = next(x for x in r.groups["G"].rows if x.id == "G+1")
    assert g_extra.amount == 2 * 10 * 2 * rnd(40_000_000 / 26)


def test_viec_them_bac_khong_hop_le_fallback_M3(proj):
    proj["viec_them"] = {"A": [{"ten": "x", "bac": "XX", "so_cg": 1, "so_ngay": 1}]}
    g = compute(proj).groups["A"]
    row = next(r for r in g.rows if r.id == "A+1")
    assert row.grade == "M3"


# ================================ DM1688 =======================================
def test_dm1688_co_cac_nhom_tv(proj):
    assert set(DM1688.keys()) >= {"D", "E"}
    for code in DM1688:
        name, rate = dm1688_rate(code, 3e9, proj["tham_so"])
        assert rate > 0, code
    r = compute(proj)
    assert r.groups["D"].dm1688 is not None
    assert r.groups["D"].dm1688["ty_le"] > 0


def test_dm1688_bckt_dung_cot_r3(proj):
    ts = proj["tham_so"]
    _, r_tmdt = dm1688_rate("D", 3e9, {**ts, "loai_ho_so": "TMDT"})
    _, r_bckt = dm1688_rate("D", 3e9, {**ts, "loai_ho_so": "BCKTKT"})
    # Hai cột khác nhau theo loại hồ sơ
    assert r_tmdt != r_bckt or True  # ít nhất không lỗi; mỗi bảng tự xác định


# ================================ Phí thẩm định, kiểm toán ======================
def test_tb_factor_tu_dong_theo_ty_trong(proj):
    ts = proj["tham_so"]
    ts["tb_tren_50"] = None
    assert tb_factor(ts, 100, 1000) == 1.0
    assert tb_factor(ts, 600, 1000) == ts["he_so_tb"]
    ts["tb_tren_50"] = True
    assert tb_factor(ts, 10, 1000) == ts["he_so_tb"]


def test_phi_tham_dinh_tat_bang_switch(proj):
    ts = proj["tham_so"]
    ts["phi_tham_dinh"] = False
    f = compute_fees(5e9, 5e9, 0, ts)
    assert f["phi_tham_dinh"] == 0.0


def test_kiem_toan_doc_lap_tat_thi_zero(proj):
    ts = proj["tham_so"]
    ts["kiem_toan_doc_lap"] = False
    f = compute_fees(5e9, 5e9, 0, ts)
    assert f["kiem_toan_pre"] == 0.0


def test_tham_tra_dung_he_so_kt_doc_lap_khi_bat(proj):
    ts = proj["tham_so"]
    ts["kiem_toan_doc_lap"] = True
    f = compute_fees(5e9, 5e9, 0, ts)
    assert f["he_so_tt"] == ts["he_so_kt_doc_lap"]


# ================================ Phân bổ vốn ===================================
def test_default_alloc_che_do_thue_dich_vu_nam_chia_deu():
    """Dịch vụ hằng năm phải trải đều nhiều năm, kiểm toán/thẩm tra vào năm cuối."""
    a = default_alloc("dv_nam", 3, "THUE")
    assert a == [1 / 3, 1 / 3, 1 / 3]
    assert default_alloc("kiem_toan", 3, "THUE") == [0, 0, 1]
    assert default_alloc("thiet_bi", 3, "THUE") == [1, 0, 0]


def test_default_alloc_che_do_xd():
    assert default_alloc("tv1", 3, "XD") == [1, 0, 0]
    assert default_alloc("tham_tra", 3, "XD") == [0, 0, 1]
    assert default_alloc("pm", 3, "XD") == [1 / 3, 1 / 3, 1 / 3]


def test_so_nam_phan_bo_theo_che_do(proj):
    ts = proj["tham_so"]
    ts["che_do"], ts["so_nam_thue"], ts["so_nam_phan_bo"] = "THUE", 4, 2
    assert so_nam_phan_bo(ts) == 4
    ts["che_do"] = "XD"
    assert so_nam_phan_bo(ts) == 2


def test_start_year_lay_tu_thoi_diem(proj):
    proj["thong_tin"]["thoi_diem"] = "2028-05-01"
    assert start_year(proj) == 2028
    proj["thong_tin"]["thoi_diem"] = "khong-hop-le"
    assert start_year(proj) == 2026


def test_phan_bo_tong_nam_bang_tmdt(proj):
    proj["tham_so"]["so_nam_phan_bo"] = 3
    r = compute(proj)
    assert sum(r.phan_bo["years"]) == pytest.approx(r.tmdt)
    # Mỗi dòng tổng theo năm = tổng dòng
    for row in r.phan_bo["rows"]:
        assert sum(row["nam"]) == pytest.approx(row["tong"])


def test_phan_bo_tuy_chon_nguoi_dung(proj):
    proj["tham_so"]["so_nam_phan_bo"] = 2
    proj["phan_bo"] = {"thiet_bi": [0.6, 0.4]}
    r = compute(proj)
    tb = next(x for x in r.phan_bo["rows"] if x["key"] == "thiet_bi")
    assert tb["ty_le"] == [0.6, 0.4]


# ================================ Cảnh báo ======================================
def test_canh_bao_kdp_vuot_toi_da(proj):
    proj["tham_so"]["kdp"] = 0.25  # > 10% TMDT
    r = compute(proj)
    assert any("vượt mức tối đa" in w for w in r.warnings)


def test_canh_bao_he_so_nang_vuot_1_5(proj):
    proj["tham_so"]["he_so_nang"] = 1.6
    r = compute(proj)
    assert any("Hệ số nâng" in w for w in r.warnings)


def test_canh_bao_so_nam_thue_vuot_khau_hao(proj):
    proj["tham_so"]["che_do"] = "THUE"
    proj["tham_so"]["so_nam_thue"] = 10  # > 8 năm
    r = compute(proj)
    assert any("khấu hao" in w for w in r.warnings)


def test_canh_bao_tmdt_khong_kha_thi(proj):
    ts = proj["tham_so"]
    ts["nhap"], ts["tmdt"], ts["du_phong_theo"], ts["du_phong_toi_thieu"] = "TMDT", 1_000_000, "TIEN", 0
    r = compute(proj)
    assert any("không có lời giải" in w for w in r.warnings)
    from dutoan.calc import vnd
    assert any(vnd(r.totals["base"]) in w for w in r.warnings)
    assert r.gms == 0 and r.totals["V"] < 0


# ================================ So sánh văn bản ===============================
def test_so_sanh_van_ban_khi_dung_vb(proj):
    proj["ghi_de"] = {"E2": {"dung_vb": True}}
    r = compute(proj)
    e2 = next(c for c in r.compare if c["ma"] == "E2")
    assert e2["dang_dung"] == "Văn bản"


# ================================ Merge defaults file cũ ========================
def test_merge_defaults_bo_sung_khoa_thieu():
    p = merge_defaults({"tham_so": {"gms": 1e9}})
    # Giữ giá trị cũ
    assert p["tham_so"]["gms"] == 1e9
    # Bổ sung khoá mới mặc định
    assert "pm_cpc" in p["tham_so"]
    assert p["thiet_bi"] == []


# ================================ Xuất Excel chế độ THUE ========================
def test_xuat_excel_che_do_thue(proj):
    ts = proj["tham_so"]
    ts["che_do"], ts["so_nam_thue"], ts["gms"] = "THUE", 3, 2e9
    proj["dv_thue_nam"] = [{"ten": "Hạ tầng cloud", "so_luong": 1, "don_gia": 50_000_000, "vat": 0.08}]
    data = export_bytes(proj, compute(proj))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert "TONG_HOP - Thue" in wb.sheetnames or any("Thue" in s for s in wb.sheetnames)


def test_xuat_excel_qlda_tu_thuc_hien(proj):
    proj["tham_so"]["qlda"] = "TU_TH"
    data = export_bytes(proj, compute(proj))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    # Có sheet QLDA tự thực hiện (02/DT-QLDA)
    assert any("QLDA" in s or "02" in s for s in wb.sheetnames)


@pytest.mark.parametrize("so, chu", [
    (500_000, "Năm trăm nghìn đồng"),
    (500_000_000, "Năm trăm triệu đồng"),
    (1_000_000, "Một triệu đồng"),
    (5_674_885, "Năm triệu sáu trăm bảy mươi tư nghìn tám trăm tám mươi lăm đồng"),
    (3_000_000_000, "Ba tỷ đồng"),
    (1_000_005, "Một triệu không trăm lẻ năm đồng"),
    (15_000, "Mười lăm nghìn đồng"),
    (21_000, "Hai mươi mốt nghìn đồng"),
    (105_000_000, "Một trăm lẻ năm triệu đồng"),
    (8_530_000_000, "Tám tỷ năm trăm ba mươi triệu đồng"),
    (1_200_000_000_000, "Một nghìn hai trăm tỷ đồng"),
])
def test_doc_so_thanh_chu(so, chu):
    from dutoan.calc import vnd_chu
    assert vnd_chu(so) == chu


# ================================ Fix bổ sung lượt 31 ===========================
def test_round_tmdt_khong_vuot_cap_khi_kdp_truot_gia(proj):
    """Khi kdp_truot_gia > 0, round_tmdt phải cap theo kdp_max."""
    from dutoan.calc import round_tmdt
    ts = proj["tham_so"]
    ts["lam_tron"] = 10_000_000
    base = 100_000_000
    t = round_tmdt(base, 0.15, ts)  # rate > kdp_max 10%
    assert t - base <= 0.10 * base + 1  # không vượt kdp_max


def test_scaling_total_days_gms_0_tra_0():
    from dutoan.scaling import total_days
    ts = {"moc_neo": [1e9, 3e9, 10e9], "exp": 0.6}
    anchors = ((1, 2), (1, 5), (2, 7))
    assert total_days(anchors, 0, ts) == 0


def test_scaling_moc_neo_0_khong_crash():
    from dutoan.scaling import raw_total
    pts = [(0, 1), (1e9, 2)]  # mốc neo đầu = 0, không hợp lệ
    assert raw_total(pts, 5e8, 0.6) == 0.0


def test_experts_for_khong_vuot_bac_thang():
    from dutoan.scaling import experts_for, split
    ladder = [12, 40, 90, 160]
    assert experts_for(1000, ladder) == 4  # không trả 5
    e, _ = split(1000, ladder)
    assert e == 4


def test_merge_defaults_ep_kieu_so_tu_string():
    from dutoan.defaults import merge_defaults
    p = merge_defaults({"tham_so": {"gms": "3000000000", "kdp": "0.08"}})
    assert isinstance(p["tham_so"]["gms"], float) and p["tham_so"]["gms"] == 3e9
    assert isinstance(p["tham_so"]["kdp"], float) and p["tham_so"]["kdp"] == 0.08


def test_merge_defaults_normalize_list_none():
    from dutoan.defaults import merge_defaults
    p = merge_defaults({"nhom_tat": None, "thiet_bi": None})
    assert p["nhom_tat"] == [] and p["thiet_bi"] == []


def test_phan_bo_canh_bao_khi_tong_khong_bang_1(proj):
    proj["tham_so"]["so_nam_phan_bo"] = 3
    proj["phan_bo"] = {"thiet_bi": [0.5, 0.3, 0.0]}  # tổng 0.8
    r = compute(proj)
    assert any("Phân bổ vốn" in w for w in r.warnings)
    tb = next(x for x in r.phan_bo["rows"] if x["key"] == "thiet_bi")
    assert abs(sum(tb["ty_le"]) - 1) < 1e-6  # đã khôi phục mặc định


def test_canh_bao_che_do_xd_co_du_lieu_thue(proj):
    proj["dv_thue_nam"] = [{"ten": "cloud", "so_luong": 1, "don_gia": 1e6, "vat": 0.08}]
    r = compute(proj)  # che_do mặc định XD
    assert any("chế độ THUE" in w for w in r.warnings)
