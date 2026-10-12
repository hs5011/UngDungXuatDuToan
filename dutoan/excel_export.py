"""Xuất Excel theo bố cục 2 template (BẢNG TMĐT + dự toán chi tiết), có công thức sống.

Tham số nằm ở sheet 01_THAM_SO dưới dạng tên (Named Range); các sheet chi tiết tham chiếu
tên đó. Số chuyên gia × số ngày (kết quả mô hình co giãn) được ghi dạng giá trị, sửa tay được
ngay trong Excel. Ở chế độ nhập Gms, TMĐT ↔ kiểm toán/lệ phí phụ thuộc vòng nên file bật
tính lặp (iterative calculation).

Mỗi ô công thức còn được ghi kèm GIÁ TRỊ TÍNH SẴN (lấy từ calc) để file xem đúng số ngay cả ở
trình xem không tự tính công thức (xem trước trên Zalo/Outlook/điện thoại, Google Drive...).
"""
import io
import re
import zipfile
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties

from .calc import (SEC_NOTES, sec_kh, sec_name, kdp_max, ks_month_price, ks_avg_month, day_price, line_pre,
                   tham_tra_tk_on)
from .catalog import GROUP_BY_CODE, DM1688, DM1688_X
from .defaults import GRADES, GRADE_NOTE, LOAI_HO_SO

NUM = "#,##0"
PCT = "0.00%"
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
SEC_FILL = PatternFill("solid", fgColor="F2F2F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

SHEET_NAMES = {
    "A": "QLDA", "B": "1. Cp_LAP BCKTKT", "C": "2. Cp_THAMTRA BCKTKT", "D": "4. Cp_GIAM SAT",
    "E": "5. Cp_TV THAU", "F": "6. Cp_TV THAM DINH THAU", "G": "6_DAO_TAO", "H": "7_TRIEN_KHAI",
    "I": "KIEM_THU_TT41", "J": "QTVH",
}

CG_NOTE = ("Căn cứ xác định đơn giá lương chuyên gia tư vấn theo Thông tư số 07/2026/TT-BNV của Bộ Nội vụ "
           "(mức tối đa; 26 ngày/tháng). Số chuyên gia × số ngày theo mô hình co giãn theo quy mô Gms.")


class Book:
    def __init__(self):
        self.wb = Workbook()
        self.wb.remove(self.wb.active)

    def sheet(self, title, index=None):
        return self.wb.create_sheet(title[:31], index)

    def name(self, name, ws, cell):
        col = "".join(c for c in cell if c.isalpha())
        row = "".join(c for c in cell if c.isdigit())
        ref = f"'{ws.title}'!${col}${row}"
        self.wb.defined_names[name] = DefinedName(name, attr_text=ref)


def ngay_vn(iso):
    """'2026-10-09' → '09/10/2026'."""
    try:
        y, m, d = str(iso)[:10].split("-")
        return f"{d}/{m}/{y}"
    except ValueError:
        return iso or ""


def ref(ws, cell):
    return f"'{ws.title}'!{cell}"


def put(ws, cell, value, fmt=None, bold=False, fill=None, border=True, align=None, cv=None):
    """Ghi ô; cv = giá trị tính sẵn của ô công thức (ghi kèm vào file khi lưu)."""
    c = ws[cell]
    c.value = value
    if cv is not None and isinstance(value, str) and value.startswith("="):
        ws.__dict__.setdefault("_cv", {})[cell] = cv
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = BOLD
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    if align:
        c.alignment = align
    return c


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def header(ws, row, labels):
    for i, lab in enumerate(labels, start=1):
        put(ws, f"{get_column_letter(i)}{row}", lab, bold=True, fill=HEAD_FILL, align=CENTER)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True)


# ====================================================================== 01_THAM_SO
def write_params(bk, proj, res):
    ts = proj["tham_so"]
    info = proj["thong_tin"]
    ws = bk.sheet("01_THAM_SO")
    title(ws, "Thông số dự án và lựa chọn trường hợp", "Ô nền vàng là ô nhập; các sheet khác tham chiếu theo tên.")
    widths(ws, [46, 22, 18, 70])
    header(ws, 3, ["Tham số", "Giá trị", "Tên ô", "Ghi chú"])
    r = [4]

    def row(label, value, nm=None, fmt=None, note=""):
        i = r[0]
        put(ws, f"A{i}", label)
        put(ws, f"B{i}", value, fmt=fmt, fill=INPUT_FILL)
        put(ws, f"C{i}", nm or "")
        put(ws, f"D{i}", note, align=WRAP)
        if nm:
            bk.name(nm, ws, f"B{i}")
        r[0] += 1
        return f"B{i}"

    def section(text):
        i = r[0]
        put(ws, f"A{i}", text, bold=True, fill=SEC_FILL)
        for c in "BCD":
            put(ws, f"{c}{i}", None, fill=SEC_FILL)
        r[0] += 1

    section("1. Thông tin chung")
    row("Tên dự án", info.get("ten_du_an") or "", "TEN_DU_AN")
    row("Chủ đầu tư", info.get("chu_dau_tu") or "", "CHU_DAU_TU")
    row("Thời điểm lập dự toán", ngay_vn(info.get("thoi_diem")))
    row("Chế độ", "Thuê dịch vụ" if ts["che_do"] == "THUE" else "Xây dựng (mua sắm)")
    row("Loại hồ sơ", LOAI_HO_SO[ts["loai_ho_so"]], note="GDP tối đa theo TT39: TMĐT ≤ 10%; dự toán ≤ 7%; BCKT-KT ≤ 5%; dự án thuê ≤ 10%.")
    row("Thiết kế", "01 bước" if ts["thiet_ke"] == 1 else "02 bước")
    row("Có phần mềm nội bộ", "Có" if ts["co_pm_noi_bo"] else "Không")
    row("Quản lý dự án", "Thuê tư vấn" if ts["qlda"] == "THUE" else "CĐT tự thực hiện (TT 70/2024)")
    row("Đầu vào", "Nhập TMĐT" if ts["nhap"] == "TMDT" else "Nhập Gms")

    section("2. Đầu vào và dự phòng")
    row("Giá trị phần mềm Gms (đồng)", res.gms, "GMS", NUM,
        "Chế độ nhập TMĐT: Gms là kết quả giải lặp (lớn nhất để dự phòng nhỏ nhất)." if ts["nhap"] == "TMDT" else "Giá trị nhập.")
    if ts["nhap"] == "TMDT":
        row("TMĐT nhập (đồng)", ts["tmdt"], "TMDT_NHAP", NUM)
    row("Kdp áp dụng", ts["kdp"] if ts["kdp"] is not None else kdp_max(ts), "KDP", PCT,
        "Nhập trong giới hạn pháp lý; không mặc định." if ts["kdp"] is not None else "CHƯA NHẬP – tạm dùng mức tối đa.")
    row("Kdp tối đa", kdp_max(ts), "KDP_MAX", PCT)
    row("Dự phòng trượt giá (NĐ 224 Điều 49.7)", ts["kdp_truot_gia"], "KDP_TG", PCT)
    if ts["nhap"] == "TMDT":
        row("Dự phòng tối thiểu (đồng)", ts["du_phong_toi_thieu"], None, NUM, "Dùng khi giải ngược Gms từ TMĐT.")
    row("Làm tròn TMĐT lên bội số (đồng)", ts["lam_tron"] or 1, "LAM_TRON", NUM)
    row("Số năm thuê", ts["so_nam_thue"], "SO_NAM")
    if ts["che_do"] != "THUE":
        row("Số năm phân bổ vốn", ts["so_nam_phan_bo"], None, None, "Xem sheet 'Bang Tong hop_Phan bo von'.")
    row("Thời gian khấu hao PM (năm)", ts["khau_hao_nam"], "KH_NAM")
    row("Thuế GTGT", ts["vat"], "VAT", PCT, "Kiểm tra hiệu lực mức thuế theo thời điểm lập.")

    section("3. Chuyên gia tư vấn (TT 07/2026/TT-BNV)")
    for g in GRADES:
        row(f"Lương tháng tối đa {g}", ts["luong"][g], f"LUONG_{g}", NUM, GRADE_NOTE[g])
    row("Hệ số nâng lương chuyên gia (≤ 1,5)", ts["he_so_nang"], "HS_NANG", "0.00")
    row("Số ngày làm việc/tháng", ts["ngay_thang"], "NGAY_THANG")
    c = ts["cql"]
    row("Cql – ngưỡng 1 (Ccg ≤)", c["nguong1"], "CQL_NG1", NUM, "TT 39 Bảng 3.2")
    row("Cql – ngưỡng 2 (Ccg <)", c["nguong2"], "CQL_NG2", NUM)
    row("Cql – tỷ lệ bậc 1", c["r1"], "CQL_R1", PCT)
    row("Cql – tỷ lệ bậc 2", c["r2"], "CQL_R2", PCT)
    row("Cql – tỷ lệ bậc 3", c["r3"], "CQL_R3", PCT)
    row("Thu nhập chịu thuế tính trước (tư vấn)", ts["tn_tv"], "TN_TV", PCT)

    section("4. Phần mềm nội bộ, dịch vụ (theo đề xuất)")
    row("Chi phí chung (CPC)", ts["pm_cpc"], "PM_CPC", PCT, "Không có tỷ lệ pháp lý – theo đề xuất.")
    row("Thu nhập chịu thuế tính trước (TNCT)", ts["pm_tn"], "PM_TN", PCT, "Theo đề xuất.")
    row("Thuế/phí phần mềm nội bộ", ts["pm_thue"], "PM_THUE", PCT, "Template: 0%.")
    row("Lương cơ sở (đồng)", ts["luong_co_so"], "LCS", NUM,
        "Dùng cho đơn giá kỹ sư (sheet Dongia_NgayCong) và lương QLDA tự thực hiện. Từ 01/07/2026 – chưa đối chiếu.")
    row("Vùng (hệ số điều chỉnh Hdc)", f"Vùng {ts['vung']} – Hdc {ts['hdc'][ts['vung']]:.2f}".replace(".", ","),
        note="I/II/III/IV = 1,2/0,9/0,7/0,5 (giá trị dùng tính nằm ở Dongia_NgayCong!C15)")
    row("Bậc kỹ sư lấy bình quân", ", ".join(f"KS-B{b}" for b in (ts["ks_bac_chon"] or [1])),
        note="Hệ số lương từng bậc sửa ở sheet Dongia_NgayCong (dòng 5, ô vàng).")
    if ts["che_do"] == "THUE":
        row("QTVH theo báo giá trọn gói", "Có" if ts["j_tron_goi"] else "Không")
        if ts["j_tron_goi"]:
            row("Nhân công QTVH trọn gói / năm (đồng)", ts["j_tron_goi_nam"], "J_TRON_GOI_NAM", NUM)

    section("5. Quy mô")
    row("Số gói thầu", res.counts["goi"], "SO_GOI")
    row("Số lớp đào tạo", res.counts["lop"], "SO_LOP")
    row("Số hệ thống tích hợp", res.counts["he_thong"], "SO_HT")
    row("Hệ số co giãn EXP", ts["exp"], None, "0.00", "Mô hình co giãn ngày công theo quy mô Gms.")
    row("Ngưỡng ngày công → 1/2/3/4 chuyên gia", " / ".join(str(x) for x in ts["bac_thang"]))
    row("Mốc neo S / M / L (Gms, đồng)", " / ".join(f"{x:,.0f}".replace(",", ".") for x in ts["moc_neo"]))
    row("Việc quy mô M xuất hiện khi Gms ≥", ts["nguong_M"], None, NUM)
    row("Việc quy mô L xuất hiện khi Gms ≥", ts["nguong_L"], None, NUM)
    row("Thẩm tra thiết kế", "Có" if tham_tra_tk_on(ts) else "Không",
        note="Tự động: tắt khi có PM nội bộ (nhà thầu lập thiết kế)." if ts["tham_tra_thiet_ke"] is None else "Chọn tay.")

    section("6. Chi phí khác theo tỷ lệ")
    row("Có kiểm toán độc lập", 1 if ts["kiem_toan_doc_lap"] else 0, "KT_ON", note="1 = có; 0 = không")
    row("Kiểm toán tối thiểu (trước thuế)", ts["kt_min"], "KT_MIN", NUM)
    row("Thẩm tra quyết toán tối thiểu", ts["tt_min"], "TT_MIN", NUM)
    row("Hệ số khi thiết bị ≥ 50%", ts["he_so_tb"], "HS_TB", PCT)
    row("Thiết bị ≥ 50% (tự động = trống; 1/0 = chọn tay)", "" if ts["tb_tren_50"] is None else int(bool(ts["tb_tren_50"])), "TB_50")
    row("Hệ số thẩm tra khi có kiểm toán độc lập", ts["he_so_kt_doc_lap"], "HS_KTDL", PCT)
    row("Có tính lệ phí thẩm định", 1 if ts["phi_tham_dinh"] else 0, "PTD_ON")
    row("Hệ số lệ phí (TT 64 giảm 50% × có thẩm tra 50%)", res.fees["he_so_ptd"], "PTD_HS", PCT)
    row("Chi phí thẩm định giá (đồng)", ts["tham_dinh_gia"], "TDG", NUM, "Theo báo giá/quyết định.")
    row("NĐ 19/2026: giám sát + đánh giá đầu tư (% chi phí QLDA)", ts["nd19_tl"], "ND19_TL", PCT, "10% + 2% + 2% + 3% (tham chiếu).")
    row("Tỷ lệ bảo hiểm QLDA tự thực hiện", ts["bh_qlda"], "BH_QLDA", PCT)

    section("7. NĐ 214/2025 Điều 14 – lựa chọn nhà thầu (tỷ lệ, tối thiểu, tối đa)")
    nd = {"khlcnt": "Lập KH tổng thể LCNT (× chi phí lập BCKT-KT)", "lap_hsmt": "Lập HSMT", "dg_hsdt": "Đánh giá HSDT",
          "td_hsmt": "Thẩm định HSMT", "td_kq": "Thẩm định KQ LCNT"}
    for k, lab in nd.items():
        rt, lo, hi = ts["nd214"][k]
        row(f"{lab} – tỷ lệ", rt, f"LCNT_{k.upper()}_R", "0.000%")
        row(f"{lab} – tối thiểu", lo, f"LCNT_{k.upper()}_MIN", NUM)
        row(f"{lab} – tối đa", hi, f"LCNT_{k.upper()}_MAX", NUM)
    return ws


# ====================================================================== đơn giá
def write_luong_cg(bk, proj):
    ts = proj["tham_so"]
    ws = bk.sheet("BANGLUONG_CG")
    title(ws, "BẢNG XÁC ĐỊNH ĐƠN GIÁ LƯƠNG CHUYÊN GIA TƯ VẤN", "Thông tư số 07/2026/TT-BNV – mức tối đa")
    widths(ws, [10, 10, 20, 18, 14, 22, 70])
    header(ws, 4, ["Mức", "Ký hiệu", "Loại đơn giá", "Giá trị", "Đơn vị", "Công thức quy đổi", "Điều kiện áp dụng"])
    r = 5
    for g in GRADES:
        put(ws, f"A{r}", f"Mức {g[1]}")
        put(ws, f"B{r}", g)
        put(ws, f"C{r}", "Lương theo tháng")
        put(ws, f"D{r}", f"=LUONG_{g}*HS_NANG", NUM, cv=ts["luong"][g] * ts["he_so_nang"])
        put(ws, f"E{r}", "đồng/tháng")
        put(ws, f"F{r}", "Mức tối đa × hệ số nâng")
        put(ws, f"G{r}", GRADE_NOTE[g], align=WRAP)
        put(ws, f"B{r + 1}", g)
        put(ws, f"C{r + 1}", "Lương theo ngày")
        put(ws, f"D{r + 1}", f"=ROUND(D{r}/NGAY_THANG,0)", NUM, cv=day_price(g, ts))
        put(ws, f"E{r + 1}", "đồng/ngày")
        put(ws, f"F{r + 1}", "Tháng / 26")
        bk.name(f"DG_{g}", ws, f"D{r + 1}")
        r += 2


def write_dongia_ks(bk, proj):
    ts = proj["tham_so"]
    ws = bk.sheet("Dongia_NgayCong")
    title(ws, "BẢNG LƯƠNG NHÂN CÔNG KỸ SƯ CNTT", "QĐ 320/QĐ-BKHCN (công thức theo template; chưa xác nhận hiệu lực sau 01/07/2026)")
    widths(ws, [6, 46] + [14] * 8)
    header(ws, 4, ["TT", "Tiêu chí"] + [f"KS-B{i}" for i in range(1, 9)])

    def vals(i):
        mlcb = (ts["ks_he_so"][i] + ts["ks_hpc"]) * ts["luong_co_so"]
        hdc = ts["hdc"][ts["vung"]]
        m = ks_month_price(i + 1, ts)
        return [None, None, ts["luong_co_so"], mlcb, hdc, mlcb * (1 + hdc), mlcb * ts["ks_bh"], m, m / ts["ngay_thang"]]
    rows = [
        ("Hệ số lương (Hcb)", lambda c, i: ts["ks_he_so"][i]),
        ("Hệ số phụ cấp (Hpc)", lambda c, i: ts["ks_hpc"]),
        ("Mức lương cơ sở (MLcs)", lambda c, i: "=LCS"),
        ("Mức lương cơ bản (MLcb)", lambda c, i: f"=({c}5+{c}6)*{c}7"),
        ("Điều chỉnh tăng thêm tiền lương (Hdc)", lambda c, i: "=HDC"),
        ("Lương (Hcb+Hpc) × MLcs × (1+Hdc)", lambda c, i: f"={c}8*(1+{c}9)"),
        ("BHXH, BHYT, BHTN người sử dụng lao động đóng", lambda c, i: f"={c}8*KS_BH"),
        ("Mức lương lao động/tháng", lambda c, i: f"={c}10+{c}11"),
        ("Mức lương lao động/ngày", lambda c, i: f"={c}12/NGAY_THANG"),
    ]
    for k, (lab, f) in enumerate(rows):
        rr = 5 + k
        put(ws, f"A{rr}", k + 1)
        put(ws, f"B{rr}", lab)
        for i in range(8):
            col = get_column_letter(3 + i)
            v = f(col, i)
            put(ws, f"{col}{rr}", v, fmt="0.00" if k in (0, 1, 4) else NUM, fill=INPUT_FILL if k in (0, 1) else None,
                cv=vals(i)[k])
        if k == 7:
            for i in range(8):
                bk.name(f"KS_THANG_{i + 1}", ws, f"{get_column_letter(3 + i)}{rr}")
    put(ws, "B15", "Hệ số điều chỉnh vùng (Hdc) áp dụng", bold=True)
    put(ws, "C15", ts["hdc"][ts["vung"]], fmt="0.00", fill=INPUT_FILL)
    put(ws, "D15", f"Vùng {ts['vung']} (I/II/III/IV = 1,2/0,9/0,7/0,5)", border=False)
    bk.name("HDC", ws, "C15")
    put(ws, "B16", "Tỷ lệ BHXH 17,5% + BHYT 3% + BHTN 1%", bold=True)
    put(ws, "C16", ts["ks_bh"], fmt=PCT, fill=INPUT_FILL)
    bk.name("KS_BH", ws, "C16")
    bacs = ts["ks_bac_chon"] or [1]
    put(ws, "B17", "Đơn giá tháng bình quân (" + "/".join(f"KS-B{b}" for b in bacs) + ")", bold=True)
    if ts["pm_don_gia_thang"]:
        put(ws, "C17", ts["pm_don_gia_thang"], NUM, fill=INPUT_FILL)
        put(ws, "D17", "Ghi đè tay", border=False)
    else:
        put(ws, "C17", "=AVERAGE(" + ",".join(f"KS_THANG_{b}" for b in bacs) + ")", NUM, cv=ks_avg_month(ts))
    put(ws, "B19", "Đổi lương cơ sở: sheet 01_THAM_SO (ô LCS). Đổi hệ số lương: dòng 5 (Hcb) của bảng trên.", border=False)
    bk.name("KS_TB_THANG", ws, "C17")


# ====================================================================== danh sách nhập
LIST_TITLES = {
    "thiet_bi": ("1_GMS_THIET_BI", "GMS - Thiết bị phần cứng"),
    "pm_thuong_mai": ("2_GMS_PM_TM", "GMS - Phần mềm thương mại/phổ biến"),
    "tb_khac": ("4_GDVTK_KHAC", "Gdvtk - Lắp đặt, tạo lập CSDL, vận chuyển, chi phí khác thuộc GTB"),
    "dv_thue_nam": ("DV_THUE_NAM", "Dịch vụ thuê hằng năm (hạ tầng, đường truyền...) – đơn giá 1 năm"),
}


def write_list(bk, proj, key):
    sname, text = LIST_TITLES[key]
    ws = bk.sheet(sname)
    title(ws, text)
    widths(ws, [8, 44, 10, 12, 16, 10, 18, 16, 18, 22, 22, 22, 22])
    header(ws, 3, ["Mã", "Khoản mục/công việc", "ĐVT", "Khối lượng", "Đơn giá trước thuế", "Thuế suất",
                   "Thành tiền trước thuế", "Thuế/phí", "Tổng cộng", "Phương pháp xác định giá",
                   "Nguồn giá/căn cứ", "Căn cứ pháp lý", "Ghi chú"])
    rows = proj[key]
    n = max(len(rows), 5)
    tot = [0.0, 0.0]
    for i in range(n):
        r = 4 + i
        d = rows[i] if i < len(rows) else {}
        pre = line_pre(d) if d else 0.0
        vat = pre * float(d.get("vat") or 0)
        tot[0] += pre
        tot[1] += vat
        put(ws, f"A{r}", d.get("ma") or (i + 1 if d else None))
        put(ws, f"B{r}", d.get("ten"), fill=INPUT_FILL, align=WRAP)
        put(ws, f"C{r}", d.get("dvt"), fill=INPUT_FILL)
        put(ws, f"D{r}", d.get("so_luong"), "#,##0.##", fill=INPUT_FILL)
        put(ws, f"E{r}", d.get("don_gia"), NUM, fill=INPUT_FILL)
        put(ws, f"F{r}", d.get("vat") if d else None, PCT, fill=INPUT_FILL)
        put(ws, f"G{r}", f"=IF(OR(D{r}=\"\",E{r}=\"\"),0,D{r}*E{r})", NUM, cv=pre)
        put(ws, f"H{r}", f"=G{r}*N(F{r})", NUM, cv=vat)
        put(ws, f"I{r}", f"=G{r}+H{r}", NUM, cv=pre + vat)
        put(ws, f"J{r}", d.get("phuong_phap"), fill=INPUT_FILL)
        put(ws, f"K{r}", d.get("nguon"), fill=INPUT_FILL)
        put(ws, f"L{r}", d.get("can_cu"), fill=INPUT_FILL)
        put(ws, f"M{r}", d.get("ghi_chu"), fill=INPUT_FILL)
    t = 4 + n
    put(ws, f"B{t}", "TỔNG", bold=True)
    for c, v in zip("GHI", (tot[0], tot[1], tot[0] + tot[1])):
        put(ws, f"{c}{t}", f"=SUM({c}4:{c}{t - 1})", NUM, bold=True, cv=v)
    return ref(ws, f"G{t}"), ref(ws, f"H{t}")


def write_gk_khac(bk, proj):
    """Chi phí khác nhập tay thuộc Gk (bước 2e)."""
    ws = bk.sheet("GK_KHAC")
    title(ws, "Chi phí khác nhập tay thuộc Gk", "Ô nền vàng là ô nhập; tổng đưa vào BẢNG TMĐT (dòng k6).")
    widths(ws, [8, 60, 20])
    header(ws, 3, ["STT", "Khoản mục", "Giá trị (đồng)"])
    rows = [d for d in proj["gk_khac"] if d.get("ten") or d.get("gia_tri")]
    n = max(len(rows), 3)
    for i in range(n):
        d = rows[i] if i < len(rows) else {}
        put(ws, f"A{4 + i}", i + 1 if d else None)
        put(ws, f"B{4 + i}", d.get("ten"), fill=INPUT_FILL)
        put(ws, f"C{4 + i}", d.get("gia_tri"), NUM, fill=INPUT_FILL)
    t = 4 + n
    put(ws, f"B{t}", "TỔNG", bold=True)
    put(ws, f"C{t}", f"=SUM(C4:C{t - 1})", NUM, bold=True, cv=sum(float(d.get("gia_tri") or 0) for d in rows))
    return ref(ws, f"C{t}")


# ====================================================================== nhóm công việc
def write_group(bk, proj, res, g):
    ts = proj["tham_so"]
    cat = GROUP_BY_CODE[g.code]
    ws = bk.sheet(SHEET_NAMES[g.code])
    service = cat.section in ("TB", "DV")
    title(ws, "BẢNG TÍNH CHI PHÍ" + ("" if service else " TƯ VẤN"), f"Công việc: {g.name}")
    widths(ws, [6, 46, 7, 10, 10, 10, 14, 16, 16, 16, 8, 50])
    header(ws, 4, ["STT", "CHI PHÍ", "BẬC", "SỐ LƯỢNG CHUYÊN GIA", "SỐ NGÀY THỰC HIỆN", "HỆ SỐ (gói/lớp/HT/năm)",
                   "ĐƠN GIÁ NGÀY CÔNG", "THÀNH TIỀN", "THEO VĂN BẢN (trước thuế)", "THEO VĂN BẢN (sau thuế)",
                   "DÙNG VB (x)", "GHI CHÚ"])
    put(ws, "A5", "I", bold=True, fill=SEC_FILL)
    put(ws, "B5", "Chi phí chuyên gia trực tiếp (Ccg)", bold=True, fill=SEC_FILL)
    put(ws, "L5", CG_NOTE, align=WRAP)
    rows = [r for r in g.rows if not r.excluded]
    first = 6
    last = first + max(len(rows), 1) - 1
    ccg = f"H5"
    mult_name = {"": None, "goi": "SO_GOI", "lop": "SO_LOP", "he_thong": "SO_HT", "nam": "SO_NAM"}
    pending_nd19 = []
    for i, row in enumerate(rows):
        r = first + i
        put(ws, f"A{r}", i + 1)
        put(ws, f"B{r}", row.name, align=WRAP)
        put(ws, f"C{r}", row.grade, fill=INPUT_FILL)
        put(ws, f"D{r}", row.experts, fill=INPUT_FILL)
        put(ws, f"E{r}", row.days, fill=INPUT_FILL)
        mn = mult_name.get(row.mult_kind)
        put(ws, f"F{r}", f"={mn}" if mn else row.mult, "#,##0.##", fill=INPUT_FILL if row.mult_kind == "tuy_chon" else None,
            cv=row.mult)
        put(ws, f"G{r}", f"=INDIRECT(\"DG_\"&C{r})", NUM, cv=row.price)
        put(ws, f"H{r}", f"=IF(LOWER(K{r})=\"x\",0,D{r}*E{r}*F{r}*G{r})", NUM, cv=row.amount)
        if row.ref.startswith("nd214"):
            k = row.ref.replace("nd214_", "").upper()
            lim = f"LCNT_{k}_R,LCNT_{k}_MIN),LCNT_{k}_MAX),0)"
            if k == "KHLCNT":
                put(ws, f"I{r}", f"=ROUND(MIN(MAX(T_TV1_PRE*{lim}", NUM, cv=row.ref_pre)
            else:
                put(ws, f"I{r}", f"=ROUND(MIN(MAX(GIA_GOI*{lim}*SO_GOI", NUM, cv=row.ref_pre)
        elif row.ref == "nd19":
            pending_nd19.append((r, row.ref_pre))
        if row.ref:
            put(ws, f"J{r}", f"=ROUND(I{r}*(1+VAT),0)", NUM, cv=row.ref_value)
        put(ws, f"K{r}", "x" if row.use_ref else None, fill=INPUT_FILL if row.ref else None, align=CENTER)
        note = ("Công việc bổ sung (nhập tay)" if row.mult_kind == "tuy_chon" else
                f"Tự động: {row.auto_experts} CG × {row.auto_days} ngày" + (" – ĐÃ GHI ĐÈ" if row.overridden else ""))
        if row.ref_note:
            note += f". {row.ref_note}"
        put(ws, f"L{r}", note, align=WRAP)
    put(ws, ccg, f"=SUM(H{first}:H{last})", NUM, bold=True, fill=SEC_FILL, cv=g.ccg)

    s = last + 1
    lines = [
        ("II", "Chi phí quản lý (Cql)" if not service else "Chi phí chung (CPC)", "rate"),
        ("III", "Chi phí khác (Ck)", "ck"),
        ("IV", "Thu nhập chịu thuế tính trước (TN)", "tn"),
        ("V", "Tổng cộng chi phí trước thuế", "pre"),
        ("VI", "Thuế giá trị gia tăng", "vat"),
        ("", "Tổng cộng sau thuế (theo chuyên gia)", "cg"),
        ("", "Cộng các dòng chọn theo văn bản – trước thuế", "vbpre"),
        ("", "Cộng các dòng chọn theo văn bản – sau thuế", "vbpost"),
        ("", "Ghi đè tổng nhóm (sau thuế, để trống nếu không dùng)", "ov"),
        ("", "TỔNG NHÓM – trước thuế", "gpre"),
        ("", "TỔNG NHÓM – thuế GTGT", "gvat"),
        ("", "TỔNG NHÓM – sau thuế", "gpost"),
    ]
    cell = {}
    for k, (stt, lab, key) in enumerate(lines):
        cell[key] = f"H{s + k}"
    rate_cell = f"G{s}"
    tn_rate = "PM_TN" if service else "TN_TV"
    rate_f = "=PM_CPC" if service else f"=IF({ccg}<=CQL_NG1,CQL_R1,IF({ccg}<CQL_NG2,CQL_R2,CQL_R3))"
    sum4 = f"{ccg}+{cell['rate']}+{cell['ck']}+{cell['tn']}"
    formulas = {
        "rate": f"={ccg}*{rate_cell}",
        "ck": float(proj.get("ck", {}).get(g.code) or 0),
        "tn": f"=ROUND({tn_rate}*({ccg}+{cell['rate']}+{cell['ck']}),0)",
        "pre": f"=ROUND({sum4},0)",
        "vat": f"=ROUND(VAT*({sum4}),0)",
        "cg": f"={cell['pre']}+{cell['vat']}",
        "vbpre": f"=SUMIF(K{first}:K{last},\"x\",I{first}:I{last})",
        "vbpost": f"=SUMIF(K{first}:K{last},\"x\",J{first}:J{last})",
        "ov": g.override if g.override is not None else None,
        "gpre": f"=IF({cell['ov']}=\"\",{cell['pre']}+{cell['vbpre']},ROUND({cell['ov']}/(1+VAT),0))",
        "gvat": f"=IF({cell['ov']}=\"\",{cell['vat']}+{cell['vbpost']}-{cell['vbpre']},{cell['ov']}-{cell['gpre']})",
        "gpost": f"={cell['gpre']}+{cell['gvat']}",
    }
    vals = {"rate": g.cql, "tn": g.tn, "pre": g.cg_pre, "vat": g.cg_vat, "cg": g.cg_total, "vbpre": g.ref_pre,
            "vbpost": g.ref_total, "gpre": g.pre, "gvat": g.vat, "gpost": g.total}
    descr = {
        "rate": "Bảng 3.2 TT 39: Ccg ≤ 1 tỷ: 55%; 1 < Ccg < 5 tỷ: 50%; Ccg ≥ 5 tỷ: 45%" if not service else "Theo đề xuất",
        "tn": f"Tỷ lệ × (Ccg + Cql + Ck)",
        "pre": "Ccg + Cql + Ck + TN",
        "vat": "VAT × (Ccg + Cql + Ck + TN)",
    }
    for k, (stt, lab, key) in enumerate(lines):
        r = s + k
        bold = key in ("cg", "gpost", "gpre", "gvat")
        put(ws, f"A{r}", stt, bold=bold)
        put(ws, f"B{r}", lab, bold=bold)
        put(ws, f"H{r}", formulas[key], NUM, bold=bold, fill=INPUT_FILL if key in ("ck", "ov") else None, cv=vals.get(key))
        if key in descr:
            put(ws, f"L{r}", descr[key], align=WRAP)
    put(ws, rate_cell, rate_f, PCT, cv=g.cql_rate)
    for r, v in pending_nd19:
        put(ws, f"I{r}", f"=ROUND(ND19_TL*({ccg}-H{r})*(1+{rate_cell})*(1+{tn_rate}),0)", NUM, cv=v)

    nm = cat.short
    bk.name(f"G_{nm}_PRE", ws, cell["gpre"])
    bk.name(f"G_{nm}_VAT", ws, cell["gvat"])

    if g.dm1688:
        r = s + len(lines) + 2
        put(ws, f"B{r}", "TÍNH THEO ĐỊNH MỨC QĐ 1688 – CHỈ THAM CHIẾU (hết hiệu lực từ 01/07/2026)", bold=True, border=False)
        header(ws, r + 1, ["", "Bảng", "", "", "", "", "Tỷ lệ %", "GT + VAT"])
        put(ws, f"B{r + 2}", g.dm1688["bang"])
        put(ws, f"G{r + 2}", g.dm1688["ty_le"], "0.000")
        put(ws, f"H{r + 2}", f"=ROUND(GMS*G{r + 2}/100*(1+VAT),0)", NUM, cv=g.dm1688["gia_tri"])
        put(ws, f"B{r + 3}", "Chênh lệch (định mức − chuyên gia)")
        put(ws, f"H{r + 3}", f"=H{r + 2}-{cell['cg']}", NUM, cv=g.dm1688["gia_tri"] - g.cg_total)
    ws.freeze_panes = "C5"
    return f"G_{nm}_PRE", f"G_{nm}_VAT"


# ====================================================================== PM nội bộ
def write_pm(bk, proj, res):
    ts = proj["tham_so"]
    ws = bk.sheet("3_GMS_PM_NOI_BO")
    title(ws, "GMS - Phần mềm nội bộ theo yêu cầu riêng",
          "Man-month suy ngược từ Gms: Gms = NC × (1 + CPC) × (1 + TNCT) × (1 + thuế); NC = MM × Σ(tỷ lệ × đơn giá tháng)")
    widths(ws, [8, 52, 8, 10, 10, 14, 16, 18, 30])
    header(ws, 4, ["Mã", "Khoản mục/công việc", "ĐVT", "Tỷ lệ", "Bậc KS", "Khối lượng (MM)", "Đơn giá tháng",
                   "Thành tiền", "Ghi chú"])
    put(ws, "A5", "PM", bold=True)
    put(ws, "B5", "Chi phí nhân công (NC)", bold=True)
    put(ws, "C5", "MM")
    pm = res.pm
    roles = pm["roles"]
    first, last = 6, 6 + len(roles) - 1
    put(ws, "D5", f"=SUM(D{first}:D{last})", PCT, cv=sum(x["ty_le"] for x in roles))
    put(ws, "F5", f"=H5/SUMPRODUCT(D{first}:D{last},G{first}:G{last})", "#,##0.00", bold=True, cv=pm["mm"])
    put(ws, "H5", "=GMS/((1+PM_CPC)*(1+PM_TN)*(1+PM_THUE))", NUM, bold=True, cv=pm["nc"])
    for i, rr in enumerate(roles):
        r = first + i
        put(ws, f"A{r}", f"PM.{i + 1:02d}")
        put(ws, f"B{r}", rr["vai_tro"])
        put(ws, f"C{r}", "MM")
        put(ws, f"D{r}", rr["ty_le"], PCT, fill=INPUT_FILL)
        b = rr["bac"]
        put(ws, f"E{r}", "TB" if b in ("TB", None, "") else f"B{b}", fill=INPUT_FILL)
        put(ws, f"F{r}", f"=$F$5*D{r}", "#,##0.00", cv=rr["mm"])
        put(ws, f"G{r}", "=KS_TB_THANG" if b in ("TB", None, "") else f"=KS_THANG_{int(b)}", NUM, cv=rr["don_gia_thang"])
        put(ws, f"H{r}", f"=F{r}*G{r}", NUM, cv=rr["thanh_tien"])
    r = last + 1
    put(ws, f"A{r}", "CPC"); put(ws, f"B{r}", "Chi phí chung (CPC × NC)"); put(ws, f"D{r}", "=PM_CPC", PCT, cv=ts["pm_cpc"])
    put(ws, f"H{r}", "=H5*PM_CPC", NUM, cv=pm["cpc"])
    put(ws, f"A{r + 1}", "TNCT"); put(ws, f"B{r + 1}", "Thu nhập chịu thuế tính trước (TNCT × (NC + CPC))")
    put(ws, f"D{r + 1}", "=PM_TN", PCT, cv=ts["pm_tn"]); put(ws, f"H{r + 1}", f"=(H5+H{r})*PM_TN", NUM, cv=pm["tn"])
    put(ws, f"A{r + 2}", "THUE"); put(ws, f"B{r + 2}", "Thuế/phí"); put(ws, f"D{r + 2}", "=PM_THUE", PCT, cv=ts["pm_thue"])
    put(ws, f"H{r + 2}", f"=(H5+H{r}+H{r + 1})*PM_THUE", NUM, cv=pm["thue"])
    put(ws, f"B{r + 3}", "TỔNG (NC + CPC + TNCT + THUẾ)", bold=True)
    put(ws, f"H{r + 3}", f"=H5+H{r}+H{r + 1}+H{r + 2}", NUM, bold=True, cv=pm["total"])
    put(ws, f"B{r + 5}", "Kiểm thử độc lập theo Điều 10 TT 41 không được gộp với kiểm thử nội bộ nhà thầu.", border=False)
    bk.name("PM_NC", ws, "H5")
    bk.name("PM_TONG", ws, f"H{r + 3}")
    bk.name("PM_THUE_TIEN", ws, f"H{r + 2}")
    return "=PM_TONG-PM_THUE_TIEN", "=PM_THUE_TIEN"


# ====================================================================== QLDA tự thực hiện
def write_qlda_tu_th(bk, proj, res):
    q = res.qlda_tu_th
    ws = bk.sheet("QLDA_TU_TH")
    title(ws, "DỰ TOÁN CHI QUẢN LÝ DỰ ÁN – CHỦ ĐẦU TƯ TỰ THỰC HIỆN",
          "TT 70/2024/TT-BTC – Mẫu 02/DT-QLDA: Bảng tính lương năm; nội dung chi theo Điều 30 NĐ 10/2021/NĐ-CP")
    widths(ws, [6, 30, 10, 12, 12, 12, 12, 12, 10, 18])
    header(ws, 4, ["STT", "Vị trí", "Số người", "Hệ số lương", "PC chức vụ", "PC khu vực", "PC khác",
                   "Tổng hệ số", "Số tháng", "Thành tiền"])
    rows = proj["qlda_nhan_su"]
    first = 5
    for i, d in enumerate(rows):
        r = first + i
        put(ws, f"A{r}", i + 1)
        put(ws, f"B{r}", d.get("vi_tri"), fill=INPUT_FILL)
        for c, k in zip("CDEFG", ("so_nguoi", "he_so_luong", "pc_chuc_vu", "pc_khu_vuc", "pc_khac")):
            put(ws, f"{c}{r}", d.get(k) or 0, "0.00", fill=INPUT_FILL)
        put(ws, f"H{r}", f"=SUM(D{r}:G{r})", "0.00", cv=q["rows"][i]["tong_he_so"])
        put(ws, f"I{r}", d.get("so_thang") or 0, fill=INPUT_FILL)
        put(ws, f"J{r}", f"=H{r}*LCS*C{r}*I{r}", NUM, cv=q["rows"][i]["thanh_tien"])
    last = first + max(len(rows), 1) - 1
    r = last + 1
    put(ws, f"B{r}", "Cộng tiền lương", bold=True)
    put(ws, f"J{r}", f"=SUM(J{first}:J{last})", NUM, bold=True, cv=q["luong"])
    put(ws, f"B{r + 1}", "BHXH, BHYT, BHTN, KPCĐ"); put(ws, f"J{r + 1}", f"=J{r}*BH_QLDA", NUM, cv=q["bh"])
    r2 = r + 3
    put(ws, f"B{r2}", "Nội dung chi khác (Điều 30 NĐ 10/2021)", bold=True)
    for i, d in enumerate(proj["qlda_chi_khac"]):
        put(ws, f"B{r2 + 1 + i}", d.get("ten"))
        put(ws, f"J{r2 + 1 + i}", d.get("gia_tri") or 0, NUM, fill=INPUT_FILL)
    e = r2 + len(proj["qlda_chi_khac"])
    t = e + 1
    put(ws, f"B{t}", "TỔNG DỰ TOÁN CHI QLDA", bold=True)
    put(ws, f"J{t}", f"=J{r}+J{r + 1}+SUM(J{r2 + 1}:J{e})", NUM, bold=True, cv=q["total"])
    bk.name("QLDA_TTH", ws, f"J{t}")
    return "=QLDA_TTH", 0


# ====================================================================== bảng định mức
def write_tables(bk, proj):
    ts = proj["tham_so"]
    ws = bk.sheet("DinhMuc_ND193_TT28")
    title(ws, "BẢNG TỶ LỆ ĐỊNH MỨC", "Nội suy tuyến tính giữa cận dưới và cận trên; ngoài bảng lấy giá trị đầu/cuối")
    widths(ws, [44, 14, 14, 14])
    put(ws, "A4", "NĐ 193/2026/NĐ-CP Điều 20 – Kiểm toán, thẩm tra quyết toán", bold=True, border=False)
    header(ws, 5, ["TMĐT (tỷ đồng)", "% thẩm tra QT", "% kiểm toán"])
    tab = ts["bang_kiem_toan"]
    for i, (x, tt, kt) in enumerate(tab):
        r = 6 + i
        put(ws, f"A{r}", x, "#,##0", fill=INPUT_FILL)
        put(ws, f"B{r}", tt, "0.000", fill=INPUT_FILL)
        put(ws, f"C{r}", kt, "0.000", fill=INPUT_FILL)
    e1 = 6 + len(tab) - 1
    rng_kt = (f"$A$6:$A${e1}", f"$B$6:$B${e1}", f"$C$6:$C${e1}")
    s2 = e1 + 3
    put(ws, f"A{s2 - 1}", "TT 28/2023/TT-BTC – Phí thẩm định dự án", bold=True, border=False)
    header(ws, s2, ["TMĐT (tỷ đồng)", "Tỷ lệ %"])
    ptab = ts["bang_phi_tham_dinh"]
    for i, (x, rt) in enumerate(ptab):
        r = s2 + 1 + i
        put(ws, f"A{r}", x, "#,##0", fill=INPUT_FILL)
        put(ws, f"B{r}", rt, "0.0000", fill=INPUT_FILL)
    e2 = s2 + len(ptab)
    rng_ptd = (f"$A${s2 + 1}:$A${e2}", f"$B${s2 + 1}:$B${e2}")
    s3 = e2 + 3
    put(ws, f"A{s3 - 1}", "QĐ 1688 – HẾT HIỆU LỰC TỪ 01/07/2026, CHỈ THAM CHIẾU (dòng PM nội bộ, CSDL)", bold=True, border=False)
    header(ws, s3, ["Bảng"] + [str(x) for x in DM1688_X])
    for i, (code, (name, r2, r3)) in enumerate(DM1688.items()):
        row = r3 if ts["loai_ho_so"] == "BCKTKT" else r2
        put(ws, f"A{s3 + 1 + i}", name)
        for j, v in enumerate(row):
            put(ws, f"{get_column_letter(2 + j)}{s3 + 1 + i}", v, "0.000")
    return ws, rng_kt, rng_ptd


def interp_formula(x, xs, ys, nd):
    """Công thức nội suy giống calc.interp: ngoài bảng lấy đầu/cuối."""
    m = f"MATCH({x},{xs},1)"
    return (f"=ROUND(IF({x}<=INDEX({xs},1),INDEX({ys},1),IF({x}>=INDEX({xs},ROWS({xs})),INDEX({ys},ROWS({xs})),"
            f"INDEX({ys},{m})-(INDEX({ys},{m})-INDEX({ys},{m}+1))/(INDEX({xs},{m}+1)-INDEX({xs},{m}))*({x}-INDEX({xs},{m})))),{nd})")


# ====================================================================== BẢNG TMĐT
def write_summary(bk, proj, res, item_f, tables):
    ts = proj["tham_so"]
    tws, rng_kt, rng_ptd = tables
    ws = bk.sheet("BẢNG TMĐT", 1)
    loai = "TỔNG MỨC ĐẦU TƯ" if ts["loai_ho_so"] == "TMDT" else "TỔNG DỰ TOÁN"
    title(ws, f"{loai} DỰ ÁN")
    info = proj["thong_tin"]
    put(ws, "A2", '=CONCATENATE("Dự án: ",TEN_DU_AN)', border=False, cv="Dự án: " + (info.get("ten_du_an") or ""))
    put(ws, "A3", '=CONCATENATE("Chủ đầu tư: ",CHU_DAU_TU)', border=False,
        cv="Chủ đầu tư: " + (info.get("chu_dau_tu") or ""))
    ws["E4"] = f"Thời điểm lập: {ngay_vn(info.get('thoi_diem'))}"
    sec_val = {sec: [sum(getattr(i, k) for i in res.items if i.sec == sec) for k in ("pre", "vat", "post")]
               for sec in ("I", "II", "III", "IV", "V")}
    f = res.fees
    ws["G4"] = "Đơn vị tính: đồng"
    widths(ws, [6, 62, 10, 26, 34, 18, 16, 18, 40])
    header(ws, 5, ["STT", "Nội dung chi phí", "Ký hiệu", "Định mức", "Cách tính", "Giá trị trước thuế",
                   "Thuế GTGT", "Giá trị sau thuế", "Ghi chú"])
    r = 6
    sec_rows = {}
    secs = ["I", "II", "III", "IV", "V"]
    cells = {}
    fee_rows = {}
    for sec in secs:
        items = [it for it in res.items if it.sec == sec]
        sr = r
        sec_rows[sec] = sr
        put(ws, f"A{sr}", sec, bold=True, fill=SEC_FILL)
        put(ws, f"B{sr}", sec_name(sec, ts), bold=True, fill=SEC_FILL)
        put(ws, f"C{sr}", sec_kh(sec, ts), bold=True, fill=SEC_FILL)
        for c in "DE":
            put(ws, f"{c}{sr}", None, fill=SEC_FILL)
        sec_key = "sec_" + sec
        sec_note = proj.get("ghi_chu", {}).get(sec_key, SEC_NOTES[sec])
        if sec in ("II", "V") and items:
            sec_note = proj.get("ghi_chu", {}).get(items[0].key, items[0].note)
        put(ws, f"I{sr}", sec_note, fill=SEC_FILL, align=WRAP)
        r += 1
        if sec in ("II", "V"):
            if items:
                cells[items[0].key] = sr
            else:
                put(ws, f"F{sr}", 0, NUM, fill=SEC_FILL)
                put(ws, f"G{sr}", 0, NUM, fill=SEC_FILL)
            r -= 1
        else:
            for n, it in enumerate(items, start=1):
                put(ws, f"A{r}", n)
                put(ws, f"B{r}", it.label, align=WRAP)
                put(ws, f"C{r}", it.kh)
                put(ws, f"I{r}", proj.get("ghi_chu", {}).get(it.key, it.note), align=WRAP)
                cells[it.key] = r
                if it.key in ("kiem_toan", "tham_tra", "phi_tham_dinh"):
                    fee_rows[it.key] = r
                r += 1
            put(ws, f"F{sr}", f"=SUM(F{sr + 1}:F{r - 1})" if items else 0, NUM, bold=True, fill=SEC_FILL,
                cv=sec_val[sec][0])
            put(ws, f"G{sr}", f"=SUM(G{sr + 1}:G{r - 1})" if items else 0, NUM, bold=True, fill=SEC_FILL,
                cv=sec_val[sec][1])
        put(ws, f"H{sr}", f"=F{sr}+G{sr}", NUM, bold=True, fill=SEC_FILL, cv=sec_val[sec][2])
        r = max(r, sr + 1)
    tot = r
    put(ws, f"A{tot}", "*", bold=True)
    put(ws, f"B{tot}", "TỔNG CỘNG", bold=True)
    put(ws, f"C{tot}", "G", bold=True)
    put(ws, f"E{tot}", "Gtb + Gqlda + Gtv + Gk + Gdp")
    for c, v in zip("FGH", (res.totals["pre"], res.totals["vat"], res.tmdt)):
        put(ws, f"{c}{tot}", "=" + "+".join(f"{c}{sec_rows[s]}" for s in secs), NUM, bold=True, cv=v)
    put(ws, f"I{tot}", proj.get("ghi_chu", {}).get("tong", "Gtb + Gqlda + Gtv + Gk + Gdp"), bold=True, align=WRAP)
    for c in "ABCDEFGHI":
        ws[f"{c}{tot}"].fill = HEAD_FILL

    # Khối TMĐT và nội suy (như ô F53–H57 template)
    b = tot + 3
    s_post = "+".join(f"H{sec_rows[s]}" for s in ("I", "II", "III", "IV"))
    put(ws, f"E{b}", "Cơ sở tính dự phòng (Gtb + Gqlda + Gtv + Gk)")
    put(ws, f"F{b}", f"={s_post}", NUM, cv=res.totals["base"])
    put(ws, f"E{b + 1}", "TMĐT xác định")
    if ts["nhap"] == "TMDT":
        put(ws, f"F{b + 1}", "=TMDT_NHAP", NUM, bold=True, cv=res.tmdt)
        put(ws, f"I{b + 1}", "TMĐT nhập; Gms giải ngược để dự phòng nhỏ nhất.")
    else:
        put(ws, f"F{b + 1}", f"=IFERROR(IF(CEILING(F{b}*(1+KDP+KDP_TG),LAM_TRON)-F{b}>KDP_MAX*F{b},FLOOR(F{b}*(1+KDP+KDP_TG),LAM_TRON),CEILING(F{b}*(1+KDP+KDP_TG),LAM_TRON)),0)", NUM, bold=True, cv=res.tmdt)
        put(ws, f"I{b + 1}", "Làm tròn lên [cơ sở × (1 + Kdp)] (xuống nếu vượt Kdp tối đa); có vòng lặp với kiểm toán/lệ phí (đã bật tính lặp).")
    bk.name("TMDT", ws, f"F{b + 1}")
    put(ws, f"E{b + 2}", "TMĐT (tỷ đồng)")
    put(ws, f"F{b + 2}", "=TMDT/1000000000", "#,##0.000", cv=res.tmdt / 1e9)
    x = f"$F${b + 2}"
    put(ws, f"E{b + 3}", "Tỷ lệ kiểm toán (%) – NĐ 193 nội suy")
    put(ws, f"F{b + 3}", interp_formula(x, ref(tws, rng_kt[0]), ref(tws, rng_kt[2]), 4), "0.0000", cv=f["ty_le_kt"])
    put(ws, f"E{b + 4}", "Tỷ lệ thẩm tra QT (%) – NĐ 193 nội suy")
    put(ws, f"F{b + 4}", interp_formula(x, ref(tws, rng_kt[0]), ref(tws, rng_kt[1]), 4), "0.0000", cv=f["ty_le_tt"])
    put(ws, f"E{b + 5}", "Tỷ lệ phí thẩm định (%) – TT 28 nội suy")
    put(ws, f"F{b + 5}", interp_formula(x, ref(tws, rng_ptd[0]), ref(tws, rng_ptd[1]), 3), "0.0000", cv=f["ty_le_ptd"])
    put(ws, f"E{b + 6}", "Hệ số thiết bị ≥ 50%")
    put(ws, f"F{b + 6}", f"=IF(TB_50=\"\",IF(IFERROR(H{sec_rows['I']}/TMDT,0)>=0.5,HS_TB,1),IF(TB_50=1,HS_TB,1))", "0.00", cv=f["he_so_tb"])
    put(ws, f"E{b + 7}", "Giá trị đề nghị kiểm toán/quyết toán (F Gtb+Gqlda+Gtv+k3+k4)")
    k3, k4 = cells["phi_tham_dinh"], cells["tham_dinh_gia"]
    put(ws, f"F{b + 7}", f"=F{sec_rows['I']}+F{sec_rows['II']}+F{sec_rows['III']}+F{k3}+F{k4}", NUM, cv=f["co_so"])
    co_so, r_kt, r_tt, r_ptd, k_tb = f"$F${b + 7}", f"$F${b + 3}", f"$F${b + 4}", f"$F${b + 5}", f"$F${b + 6}"
    put(ws, f"E{b + 8}", "Dự phòng / cơ sở")
    put(ws, f"F{b + 8}", f"=F{sec_rows['V']}/F{b}", PCT, cv=res.totals["dp_ratio"])

    # Công thức từng dòng
    fee_f = {
        "kiem_toan": (f"=IF(KT_ON=1,MAX(ROUND({co_so}*{r_kt}/100*{k_tb},0),KT_MIN),0)", "=ROUND(F{r}*VAT,0)",
                      "Giá trị đề nghị kiểm toán × định mức × hệ số TB", f"={r_kt}/100"),
        "tham_tra": (f"=MAX(ROUND({co_so}*{r_tt}/100*{k_tb}*IF(KT_ON=1,HS_KTDL,1),0),TT_MIN)", 0,
                     "Giá trị quyết toán × định mức × hệ số TB × 50% (nếu có kiểm toán)", f"={r_tt}/100"),
        "phi_tham_dinh": (f"=IF(PTD_ON=1,ROUND(TMDT*{r_ptd}/100*PTD_HS,0),0)", 0,
                          "TMĐT × mức thu × 50% (TT 64) × 50% (có thẩm tra)", f"={r_ptd}/100"),
        "tham_dinh_gia": ("=TDG", 0, "Theo báo giá", None),
    }
    fee_rate = {"kiem_toan": f["ty_le_kt"], "tham_tra": f["ty_le_tt"], "phi_tham_dinh": f["ty_le_ptd"]}
    for it in res.items:
        rr = cells[it.key]
        if it.key in fee_f:
            fp, fv, how, dm = fee_f[it.key]
            put(ws, f"F{rr}", fp, NUM, cv=it.pre)
            put(ws, f"G{rr}", fv.format(r=rr) if isinstance(fv, str) else fv, NUM, cv=it.vat)
            put(ws, f"E{rr}", how, align=WRAP)
            if dm:
                put(ws, f"D{rr}", dm, "0.0000%", cv=fee_rate[it.key] / 100)
        elif it.key == "du_phong":
            put(ws, f"F{rr}", f"=TMDT-({s_post})", NUM, bold=True, fill=SEC_FILL, cv=it.pre)
            put(ws, f"G{rr}", 0, NUM, fill=SEC_FILL)
            put(ws, f"E{rr}", "(Gtb + Gqlda + Gtv + Gk) × Kdp", fill=SEC_FILL)
        elif it.key == "qlda":
            fp, fv = item_f[it.key]
            put(ws, f"F{rr}", fp, NUM, bold=True, fill=SEC_FILL, cv=it.pre)
            put(ws, f"G{rr}", fv, NUM, bold=True, fill=SEC_FILL, cv=it.vat)
            put(ws, f"D{rr}", "Theo chuyên gia" if ts["qlda"] == "THUE" else "Dự toán chi (TT 70)", fill=SEC_FILL)
        else:
            fp, fv = item_f[it.key]
            put(ws, f"F{rr}", fp, NUM, cv=it.pre)
            put(ws, f"G{rr}", fv, NUM, cv=it.vat)
            if it.key.startswith("tv") or it.key in ("dao_tao", "trien_khai", "kiem_thu", "qtvh"):
                put(ws, f"D{rr}", "Phương pháp chuyên gia" if it.key.startswith("tv") else "Theo khối lượng, đơn giá", align=WRAP)
        put(ws, f"H{rr}", f"=F{rr}+G{rr}", NUM, bold=it.sec in ("II", "V"), fill=SEC_FILL if it.sec in ("II", "V") else None,
            cv=it.post)
        bk.name(f"T_{it.key.upper()}_POST", ws, f"H{rr}")
        if it.key == "tv1":
            bk.name("T_TV1_PRE", ws, f"F{rr}")
    if "tv1" not in cells:
        put(ws, f"J{tot}", 0, border=False)
        bk.name("T_TV1_PRE", ws, f"J{tot}")
    put(ws, f"E{b + 9}", "Giá gói thầu (Gtb trước thuế ÷ số gói)")
    put(ws, f"F{b + 9}", f"=F{sec_rows['I']}/SO_GOI", NUM, cv=sec_val["I"][0] / res.counts["goi"] if res.counts["goi"] else 0)
    bk.name("GIA_GOI", ws, f"F{b + 9}")

    # Bảng tổng hợp
    t = b + 12
    put(ws, f"B{t}", f"BẢNG TỔNG HỢP {loai}", bold=True, border=False)
    header(ws, t + 1, ["", "Các khoản mục chi phí", "", "", "", "Trước thuế", "Thuế GTGT", "Sau thuế"])
    for i, s in enumerate(secs):
        rr = t + 2 + i
        put(ws, f"A{rr}", s)
        put(ws, f"B{rr}", sec_name(s, ts).capitalize())
        for c, v in zip("FGH", sec_val[s]):
            put(ws, f"{c}{rr}", f"={c}{sec_rows[s]}", NUM, cv=v)
    rr = t + 2 + len(secs)
    put(ws, f"B{rr}", "Tổng vốn", bold=True)
    for c, v in zip("FGH", (res.totals["pre"], res.totals["vat"], res.tmdt)):
        put(ws, f"{c}{rr}", f"=SUM({c}{t + 2}:{c}{rr - 1})", NUM, bold=True, cv=v)
    ws.freeze_panes = "C6"
    return ws, cells


# ====================================================================== phân bổ vốn
def write_phan_bo(bk, proj, res, cells):
    pb = res.phan_bo
    n = pb["n"]
    ws = bk.sheet("Bang Tong hop_Phan bo von")
    title(ws, "BẢNG TỔNG HỢP – PHÂN BỔ VỐN THEO NĂM", "Tỷ lệ phân bổ (cột bên phải) sửa được; số tiền = giá trị × tỷ lệ")
    widths(ws, [6, 60, 18] + [16] * n + [4] + [10] * n + [16])
    ycols = [get_column_letter(4 + i) for i in range(n)]
    pcols = [get_column_letter(5 + n + i) for i in range(n)]
    chk = get_column_letter(5 + 2 * n)
    header(ws, 4, ["STT", "Nội dung chi phí", "Giá trị sau thuế"] + pb["labels"] + [""] +
           [f"% {l}" for l in pb["labels"]] + ["Còn lại"])
    tms = "'BẢNG TMĐT'"
    r = 5
    sec_rows = []
    for sec in ("I", "II", "III", "IV", "V"):
        rows = [x for x in pb["rows"] if x["nhom"] == sec]
        if not rows:
            continue
        sr = r
        sec_rows.append(sr)
        put(ws, f"A{sr}", sec, bold=True, fill=SEC_FILL)
        put(ws, f"B{sr}", sec_name(sec, proj["tham_so"]), bold=True, fill=SEC_FILL)
        r += 1
        for k, x in enumerate(rows, start=1):
            put(ws, f"A{r}", k)
            put(ws, f"B{r}", x["noi_dung"], align=WRAP)
            put(ws, f"C{r}", f"={tms}!H{cells[x['key']]}", NUM, cv=x["tong"])
            for i in range(n):
                put(ws, f"{pcols[i]}{r}", x["ty_le"][i], "0.00%", fill=INPUT_FILL)
                put(ws, f"{ycols[i]}{r}", f"=C{r}*{pcols[i]}{r}", NUM, cv=x["nam"][i])
            put(ws, f"{chk}{r}", f"=C{r}-SUM({ycols[0]}{r}:{ycols[-1]}{r})", NUM, cv=x["tong"] - sum(x["nam"]))
            r += 1
        tong = sum(x["tong"] for x in rows)
        nam = [sum(x["nam"][i] for x in rows) for i in range(n)]
        for c, v in zip(["C"] + ycols + [chk], [tong] + nam + [tong - sum(nam)]):
            put(ws, f"{c}{sr}", f"=SUM({c}{sr + 1}:{c}{r - 1})", NUM, bold=True, fill=SEC_FILL, cv=v)
    put(ws, f"A{r}", "*", bold=True, fill=HEAD_FILL)
    put(ws, f"B{r}", "TỔNG CỘNG", bold=True, fill=HEAD_FILL)
    for c, v in zip(["C"] + ycols + [chk], [res.tmdt] + pb["years"] + [res.tmdt - sum(pb["years"])]):
        put(ws, f"{c}{r}", "=" + "+".join(f"{c}{x}" for x in sec_rows), NUM, bold=True, fill=HEAD_FILL, cv=v)
    put(ws, f"B{r + 2}", "Cột 'Còn lại' = giá trị − tổng các năm; khác 0 nghĩa là tỷ lệ của dòng đó chưa đủ 100%.", border=False)
    ws.freeze_panes = "C5"


# ====================================================================== tổng hợp GTB / GDV
def write_tong_hop(bk, proj, res, cells):
    ts = proj["tham_so"]
    thue = ts["che_do"] == "THUE"
    ws = bk.sheet("TONG_HOP - Thue" if thue else "TONG_HOP - XD")
    title(ws, "Tổng hợp chi phí thuê theo yêu cầu riêng" if thue else "Tổng hợp CHI PHÍ TRANG THIẾT BỊ (GTB)")
    widths(ws, [8, 70, 12, 18, 60])
    header(ws, 4, ["Khoản", "Hạng mục", "Mã", "Tổng", "Ghi chú"])
    r = 5
    for k, it in enumerate([i for i in res.items if i.sec == "I"], start=1):
        put(ws, f"A{r}", k)
        put(ws, f"B{r}", it.label, align=WRAP)
        put(ws, f"C{r}", it.kh)
        put(ws, f"D{r}", f"='BẢNG TMĐT'!H{cells[it.key]}", NUM, cv=it.post)
        put(ws, f"E{r}", proj.get("ghi_chu", {}).get(it.key, it.note), align=WRAP)
        r += 1
    put(ws, f"B{r}", "TỔNG CỘNG", bold=True)
    put(ws, f"D{r}", f"=SUM(D5:D{r - 1})", NUM, bold=True, cv=sum(i.post for i in res.items if i.sec == "I"))
    if thue:
        put(ws, f"B{r + 2}", "Tổng giá trị phần mềm nội bộ (Gms)")
        put(ws, f"D{r + 2}", "=GMS", NUM, cv=res.gms)
        put(ws, f"B{r + 3}", "Số năm theo kế hoạch thuê")
        put(ws, f"D{r + 3}", "=SO_NAM", cv=ts["so_nam_thue"])
        put(ws, f"E{r + 2}", "Thuê PM = Gms ÷ thời gian khấu hao (8 năm) × số năm thuê", align=WRAP)


# ====================================================================== so sánh chuyên gia – văn bản (6b)
def write_so_sanh(bk, res):
    ws = bk.sheet("SO_SANH_VAN_BAN")
    title(ws, "SO SÁNH SỐ TÍNH THEO CHUYÊN GIA VỚI SỐ THEO VĂN BẢN",
          "Giá trị tại thời điểm xuất (NĐ 214/2025, NĐ 19/2026); đã gồm VAT. Chênh lệch = (chuyên gia − văn bản) ÷ văn bản.")
    widths(ws, [10, 34, 50, 20, 20, 12, 14, 70])
    header(ws, 4, ["Mã việc", "Nhóm chi phí", "Công việc", "Theo chuyên gia (đ)", "Theo văn bản (đ)", "Chênh lệch",
                   "Đang dùng", "Căn cứ văn bản"])
    for i, c in enumerate(res.compare):
        r = 5 + i
        put(ws, f"A{r}", c["ma"])
        put(ws, f"B{r}", f"{c['nhom']}. {c['nhom_ten']}")
        put(ws, f"C{r}", c["cong_viec"])
        put(ws, f"D{r}", c["chuyen_gia"], NUM)
        put(ws, f"E{r}", c["van_ban"], NUM)
        put(ws, f"F{r}", f"=IF(E{r}=0,0,(D{r}-E{r})/E{r})", "+0%;-0%;0%", cv=c["lech"])
        put(ws, f"G{r}", c["dang_dung"])
        put(ws, f"H{r}", c["can_cu"])
    ws.freeze_panes = "A5"


def _fill_cached(data, wb):
    """Ghi giá trị tính sẵn vào các ô công thức (openpyxl chỉ ghi công thức, để trống <v>)."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    paths = {f"xl/worksheets/sheet{i + 1}.xml": ws.__dict__.get("_cv", {}) for i, ws in enumerate(wb.worksheets)}
    pat = re.compile(r'<c r="([A-Z]+[0-9]+)"([^>]*)><f>([^<]*)</f><v\s*/?>(?:</v>)?</c>')

    def sub(cache):
        def f(m):
            v = cache.get(m.group(1))
            if v is None:
                return m.group(0)
            if isinstance(v, str):
                return f'<c r="{m.group(1)}"{m.group(2)} t="str"><f>{m.group(3)}</f><v>{escape(v)}</v></c>'
            return f'<c r="{m.group(1)}"{m.group(2)}><f>{m.group(3)}</f><v>{float(v)!r}</v></c>'
        return f

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            buf = zin.read(item.filename)
            cache = paths.get(item.filename)
            if cache:
                buf = pat.sub(sub(cache), buf.decode("utf-8")).encode("utf-8")
            zout.writestr(item, buf)
    return out.getvalue()


# ====================================================================== điểm vào
def build_workbook(proj, res):
    ts = proj["tham_so"]
    bk = Book()
    write_params(bk, proj, res)
    write_luong_cg(bk, proj)
    write_dongia_ks(bk, proj)

    item_f = {}
    list_refs = {}
    for k in ("thiet_bi", "pm_thuong_mai", "tb_khac") + (("dv_thue_nam",) if ts["che_do"] == "THUE" else ()):
        list_refs[k] = write_list(bk, proj, k)
    item_f["thiet_bi"] = ("=" + list_refs["thiet_bi"][0], "=" + list_refs["thiet_bi"][1])
    item_f["pm_tm"] = ("=" + list_refs["pm_thuong_mai"][0], "=" + list_refs["pm_thuong_mai"][1])
    item_f["tb_khac"] = ("=" + list_refs["tb_khac"][0], "=" + list_refs["tb_khac"][1])
    if "dv_thue_nam" in list_refs:
        item_f["dv_nam"] = (f"={list_refs['dv_thue_nam'][0]}*SO_NAM", f"={list_refs['dv_thue_nam'][1]}*SO_NAM")

    pm_pre, pm_vat = write_pm(bk, proj, res)
    if ts["che_do"] == "THUE":
        item_f["thue_pm"] = ("=(PM_TONG-PM_THUE_TIEN)/KH_NAM*SO_NAM", "=PM_THUE_TIEN/KH_NAM*SO_NAM")
        if ts["j_tron_goi"]:
            nc = f"{ts['j_tron_goi_nam']:.0f}*SO_NAM"
            pre = f"{nc}*(1+PM_CPC)*(1+PM_TN)"
            item_f["qtvh"] = (f"={pre}", f"={pre}*VAT")
    else:
        item_f["pm"] = (pm_pre, pm_vat)

    key_of = {"A": "qlda", "B": "tv1", "C": "tv2", "D": "tv3", "E": "tv4", "F": "tv5",
              "G": "dao_tao", "H": "trien_khai", "I": "kiem_thu", "J": "qtvh"}
    for code in "GHJABCDEFI":
        g = res.groups.get(code)
        if g:
            pn, vn = write_group(bk, proj, res, g)
            item_f[key_of[code]] = (f"={pn}", f"={vn}")
    if res.qlda_tu_th:
        item_f["qlda"] = write_qlda_tu_th(bk, proj, res)
    item_f["gk_khac"] = ("=" + write_gk_khac(bk, proj), 0)

    tables = write_tables(bk, proj)
    _, cells = write_summary(bk, proj, res, item_f, tables)
    write_tong_hop(bk, proj, res, cells)
    write_phan_bo(bk, proj, res, cells)
    if res.compare:
        write_so_sanh(bk, res)

    bk.wb.calculation = CalcProperties(fullCalcOnLoad=True, iterate=True, iterateCount=200, iterateDelta=0.5)
    bk.wb.active = 1
    return bk.wb


def export_bytes(proj, res):
    wb = build_workbook(proj, res)
    buf = io.BytesIO()
    wb.save(buf)
    return _fill_cached(buf.getvalue(), wb)
